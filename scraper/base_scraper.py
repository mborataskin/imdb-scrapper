from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from openpyxl import load_workbook
import pandas as pd
import os

class BaseScraper:
    def __init__(self):
        chrome_options = webdriver.ChromeOptions()
        chrome_options.add_argument("--headless=new")
        chrome_options.add_argument("--window-size=1920,1080")
        chrome_options.add_argument("--disable-extensions")
        chrome_options.add_argument("enable-automation")
        chrome_options.add_argument("--disable-blink-features=AutomationControlled")
        chrome_options.add_argument("user-agent=Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/134.0.0.0 Safari/537.36")

        try:
            self.driver = webdriver.Chrome(options=chrome_options)
        except Exception as e:
            print(f"An error occurred while starting the browser: {e}")
            self.driver = None

        if not os.path.exists("data/imdb_data.xlsx"):
            try:
                pd.DataFrame().to_excel("data/imdb_data.xlsx", sheet_name="Placeholder", index=False)
            except Exception as e:
                print(f"Failed to create the Excel file: {e}")

    def remove_placeholder_sheet(self):
        try:
            workbook = load_workbook("data/imdb_data.xlsx")
            if "Placeholder" in workbook.sheetnames:
                workbook.remove(workbook["Placeholder"])
                workbook.save("data/imdb_data.xlsx")
                print("Placeholder sheet successfully removed!")
        except Exception as e:
            print(f"Failed to remove the placeholder sheet: {e}")

    def quit(self):
        if self.driver:
            self.driver.quit()