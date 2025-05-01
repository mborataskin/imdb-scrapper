from scraper.base_scraper import BaseScraper
from selenium.webdriver.common.by import By
from openpyxl import load_workbook
from openpyxl.styles import Alignment
import pandas as pd

class UpcomingMoviesScraper(BaseScraper):
    def scrape(self):
        try:
            self.driver.get("https://www.imdb.com/calendar/")
            calendar_section = self.driver.find_elements(By.CSS_SELECTOR, "article[data-testid='calendar-section']")

            upcoming_movies_data = []

            for calendar_element in calendar_section[:3]:
                try:
                    date = calendar_element.find_element(By.CSS_SELECTOR, "h3").text
                except:
                    date = "Unknown Date"

                movie_sections = calendar_element.find_elements(By.CSS_SELECTOR, ".ipc-metadata-list-summary-item__tc")

                for movie in movie_sections:
                    try:
                        title = movie.find_element(By.CSS_SELECTOR, ".ipc-metadata-list-summary-item__t").text
                        title = title.split(" (")[0]
                    except:
                        title = "N/A"

                    try:
                        genres = movie.find_elements(By.CSS_SELECTOR, "ul .ipc-inline-list__item")
                        genre_list = [genre.text for genre in genres[:3]]
                    except:
                        genre_list = []

                    try:
                        stars = movie.find_elements(By.CSS_SELECTOR, "ul .ipc-inline-list.ipc-inline-list--no-wrap.ipc-inline-list--inline.ipc-metadata-list-summary-item__stl.base .ipc-inline-list__item")
                        star_list = [star.text for star in stars[:4]]
                    except:
                        star_list = []

                    upcoming_movies_data.append({
                        "Date": date,
                        "Title": title,
                        "Genres": ", ".join(genre_list),
                        "Stars": ", ".join(star_list)
                    })

            df = pd.DataFrame(upcoming_movies_data)
            df.to_csv("data/upcoming_movies_data.csv", index=False)

            try:
                with pd.ExcelWriter("data/imdb_data.xlsx", engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                    df.to_excel(writer, sheet_name="Upcoming Movies", index=False)
            except Exception as e:
                print(f"Error while writing to Excel file: {e}")

            try:
                workbook = load_workbook("data/imdb_data.xlsx")
                sheet = workbook["Upcoming Movies"]
                previous_date = None
                start_row = 2

                for row in range(2, len(upcoming_movies_data) + 2):
                    current_date = sheet[f"A{row}"].value
                    if current_date == previous_date:
                        end_row = row
                    else:
                        if previous_date and start_row < end_row:
                            sheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
                            sheet.cell(row=start_row, column=1).alignment = Alignment(horizontal="center", vertical="center")
                        start_row = row
                        end_row = row
                    previous_date = current_date

                if start_row < end_row:
                    sheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
                    sheet.cell(row=start_row, column=1).alignment = Alignment(horizontal="center", vertical="center")

                workbook.save("data/imdb_data.xlsx")
            except Exception as e:
                print(f"Error while merging cells: {e}")

            print("Upcoming Movies saved successfully.")

        except Exception as e:
            print(f"UpcomingMoviesScraper error: {e}")